import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_simple_format_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Clean classic colors matching the sample PDF
    BLACK = RGBColor(0, 0, 0)
    DARK_GRAY = RGBColor(40, 40, 40)
    RED = RGBColor(180, 0, 0)
    BLUE_HEADER = RGBColor(70, 130, 180)
    LIGHT_BLUE = RGBColor(230, 240, 250)
    WHITE = RGBColor(255, 255, 255)

    def add_common_header(slide, title_text):
        # Title text exactly like the sample PDF
        tbox = slide.shapes.add_textbox(Inches(1.0), Inches(0.6), Inches(11.333), Inches(1.0))
        tf = tbox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.name = "Arial"
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = BLACK
        p.alignment = PP_ALIGN.CENTER

    # ==========================================
    # SLIDE 1: EXACT TITLE SLIDE (Matches PDF Page 1)
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)

    # Top college name
    top_box = s1.shapes.add_textbox(Inches(1.0), Inches(0.5), Inches(11.333), Inches(1.5))
    ttf = top_box.text_frame
    ttf.word_wrap = True
    
    p0 = ttf.paragraphs[0]
    p0.text = "Galgotias College of Engineering & Technology,\nGreater Noida"
    p0.font.name = "Times New Roman"
    p0.font.size = Pt(22)
    p0.font.bold = True
    p0.font.color.rgb = BLACK
    p0.alignment = PP_ALIGN.CENTER

    p1 = ttf.add_paragraph()
    p1.text = "DEPARTMENT OF ELECTRONIC AND COMMUNICATION ENGINEERING"
    p1.font.name = "Times New Roman"
    p1.font.size = Pt(14)
    p1.font.bold = True
    p1.space_before = Pt(8)
    p1.font.color.rgb = BLACK
    p1.alignment = PP_ALIGN.CENTER

    # Project Title in Center
    center_box = s1.shapes.add_textbox(Inches(1.0), Inches(2.8), Inches(11.333), Inches(1.5))
    ctf = center_box.text_frame
    ctf.word_wrap = True
    cp = ctf.paragraphs[0]
    cp.text = "Smart Weather Station"
    cp.font.name = "Arial"
    cp.font.size = Pt(36)
    cp.font.bold = True
    cp.font.color.rgb = BLACK
    cp.alignment = PP_ALIGN.CENTER

    cp_sub = ctf.add_paragraph()
    cp_sub.text = "Using ESP32, DHT22 Sensor and Cloud Dashboard"
    cp_sub.font.name = "Arial"
    cp_sub.font.size = Pt(18)
    cp_sub.space_before = Pt(8)
    cp_sub.font.color.rgb = DARK_GRAY
    cp_sub.alignment = PP_ALIGN.CENTER

    # Mentor on Bottom Left (Red like sample)
    m_box = s1.shapes.add_textbox(Inches(1.0), Inches(5.3), Inches(4.5), Inches(1.8))
    mtf = m_box.text_frame
    mp = mtf.paragraphs[0]
    mp.text = "MENTOR:"
    mp.font.name = "Arial"
    mp.font.size = Pt(16)
    mp.font.bold = True
    mp.font.color.rgb = RED

    mp_name = mtf.add_paragraph()
    mp_name.text = "[Mentor Name]"
    mp_name.font.name = "Arial"
    mp_name.font.size = Pt(15)
    mp_name.space_before = Pt(4)
    mp_name.font.color.rgb = BLACK

    # Presented By on Bottom Right
    p_box = s1.shapes.add_textbox(Inches(7.8), Inches(5.3), Inches(4.5), Inches(1.8))
    ptf = p_box.text_frame
    pp = ptf.paragraphs[0]
    pp.text = "PRESENTED BY:"
    pp.font.name = "Arial"
    pp.font.size = Pt(16)
    pp.font.bold = True
    pp.font.color.rgb = BLACK

    students = [
        "1. Prasoon (Roll No)",
        "2. Name (Roll No)",
        "3. Name (Roll No)"
    ]
    for s in students:
        p_s = ptf.add_paragraph()
        p_s.text = s
        p_s.font.name = "Arial"
        p_s.font.size = Pt(15)
        p_s.space_before = Pt(3)
        p_s.font.color.rgb = BLACK

    # ==========================================
    # SLIDE 2: CONTENTS (Matches PDF Page 2)
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    add_common_header(s2, "CONTENTS")

    c_box = s2.shapes.add_textbox(Inches(1.5), Inches(2.0), Inches(10.0), Inches(4.8))
    ctf = c_box.text_frame
    ctf.word_wrap = True

    items = [
        "❖  Introduction",
        "❖  Motivation",
        "❖  Objectives",
        "❖  Description",
        "❖  Conclusion",
        "❖  References"
    ]
    for i, itm in enumerate(items):
        p = ctf.paragraphs[0] if i == 0 else ctf.add_paragraph()
        p.text = itm
        p.font.name = "Arial"
        p.font.size = Pt(20)
        p.font.bold = True
        p.space_before = Pt(22)
        p.font.color.rgb = BLACK

    # ==========================================
    # SLIDE 3: INTRODUCTION (Matches PDF Page 3 - Easy Words)
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    add_common_header(s3, "INTRODUCTION")

    box3 = s3.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(11.0), Inches(5.2))
    tf3 = box3.text_frame
    tf3.word_wrap = True

    intro_points = [
        "Weather conditions like temperature and humidity affect our daily lives, health, and work.",
        "Traditional weather stations are big, expensive, and cannot be checked easily from home.",
        "Internet of Things (IoT) allows us to connect small sensors to the internet to collect weather data automatically.",
        "In this project, we built a low-cost smart weather station using an ESP32 board and a DHT22 sensor.",
        "The system reads live temperature and humidity, shows them on a small screen, and sends the data to a website so anyone can view it anytime."
    ]

    for i, pt in enumerate(intro_points):
        p = tf3.paragraphs[0] if i == 0 else tf3.add_paragraph()
        p.text = "•  " + pt
        p.font.name = "Arial"
        p.font.size = Pt(17)
        p.space_before = Pt(20)
        p.font.color.rgb = BLACK

    # ==========================================
    # SLIDE 4: MOTIVATION (Easy Words)
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    add_common_header(s4, "MOTIVATION")

    box4 = s4.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(11.0), Inches(5.2))
    tf4 = box4.text_frame
    tf4.word_wrap = True

    mot_points = [
        "Need for local weather data: City weather reports are taken far away (like airports) and do not show the exact weather in our room, office, or farm.",
        "Low cost: Commercial weather devices cost a lot of money, but this project can be made at a very low cost using simple parts.",
        "Real-time monitoring: Instead of writing down values by hand, the system updates automatically every few seconds.",
        "Access from anywhere: Anyone with a phone or computer can open our website and check the live weather and past 24-hour charts."
    ]

    for i, pt in enumerate(mot_points):
        p = tf4.paragraphs[0] if i == 0 else tf4.add_paragraph()
        p.text = "•  " + pt
        p.font.name = "Arial"
        p.font.size = Pt(17)
        p.space_before = Pt(22)
        p.font.color.rgb = BLACK

    # ==========================================
    # SLIDE 5: OBJECTIVE:- (Matches PDF Page 4 - Easy Words)
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    add_common_header(s5, "OBJECTIVE:-")

    box5 = s5.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(11.0), Inches(5.2))
    tf5 = box5.text_frame
    tf5.word_wrap = True

    objectives = [
        "To build a simple and working smart weather station using ESP32 microcontroller.",
        "To accurately measure room temperature and humidity using the DHT22 sensor.",
        "To display the current readings on a small OLED screen placed directly on the board.",
        "To calculate Heat Index (how hot it feels) and Dew Point to check human comfort.",
        "To send all sensor data to the ThingSpeak cloud platform over Wi-Fi every 20 seconds.",
        "To show live readings and a 24-hour history graph on an online website hosted on Vercel."
    ]

    for i, pt in enumerate(objectives):
        p = tf5.paragraphs[0] if i == 0 else tf5.add_paragraph()
        p.text = "•  " + pt
        p.font.name = "Arial"
        p.font.size = Pt(17)
        p.space_before = Pt(18)
        p.font.color.rgb = BLACK

    # ==========================================
    # SLIDE 6: DESCRIPTION - HARDWARE & HOW IT WORKS (Matches PDF Page 5)
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    add_common_header(s6, "Description")

    box6 = s6.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(11.0), Inches(5.2))
    tf6 = box6.text_frame
    tf6.word_wrap = True

    desc_points = [
        "Hardware components used:",
        "    1. ESP32 Microcontroller: The main brain with built-in Wi-Fi to collect and send data.",
        "    2. DHT22 Sensor: Measures temperature (from -40°C to 80°C) and humidity (0 to 100%).",
        "    3. SSD1306 OLED Display: Small 0.96-inch screen to see the values on the device.",
        "",
        "Wiring and Connections:",
        "    • DHT22 data pin is connected to ESP32 Pin GPIO 4.",
        "    • OLED display is connected using I2C wires (SDA to GPIO 21, SCL to GPIO 22).",
        "    • 3.3V and GND power pins are shared between the sensor and screen.",
        "",
        "Working Steps:",
        "    Sensor reads data  ➔  ESP32 calculates comfort values  ➔  Shows on OLED  ➔  Sends to ThingSpeak cloud via Wi-Fi."
    ]

    for i, pt in enumerate(desc_points):
        p = tf6.paragraphs[0] if i == 0 else tf6.add_paragraph()
        p.text = pt
        p.font.name = "Arial"
        p.font.size = Pt(15 if not pt.endswith(":") else 16)
        p.font.bold = pt.endswith(":")
        p.space_before = Pt(4 if not pt.endswith(":") else 12)
        p.font.color.rgb = BLACK

    # ==========================================
    # SLIDE 7: DESCRIPTION - RESULTS & WEBSITE (Matches PDF Page 6)
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    add_common_header(s7, "Description")

    # Sub-heading
    sub_box = s7.shapes.add_textbox(Inches(1.2), Inches(1.5), Inches(11.0), Inches(0.6))
    sub_p = sub_box.text_frame.paragraphs[0]
    sub_p.text = "Results and Live Cloud Data:"
    sub_p.font.name = "Arial"
    sub_p.font.size = Pt(18)
    sub_p.font.bold = True
    sub_p.font.color.rgb = BLACK

    # Table for Live Results
    rows = 5
    cols = 3
    table_shape = s7.shapes.add_table(rows, cols, Inches(1.2), Inches(2.2), Inches(11.0), Inches(2.2))
    tbl = table_shape.table
    tbl.columns[0].width = Inches(3.5)
    tbl.columns[1].width = Inches(3.5)
    tbl.columns[2].width = Inches(4.0)

    res_headers = ["Parameter", "Live Measured Value", "Meaning"]
    for j, h in enumerate(res_headers):
        cell = tbl.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = BLUE_HEADER
        cp = cell.text_frame.paragraphs[0]
        cp.text = h
        cp.font.name = "Arial"
        cp.font.size = Pt(14)
        cp.font.bold = True
        cp.font.color.rgb = WHITE

    table_data = [
        ["Temperature", "26.8 °C", "Current room temperature"],
        ["Humidity", "69.8 %", "Current moisture in air"],
        ["Heat Index", "28.5 °C", "Real feel temperature"],
        ["Dew Point", "20.8 °C", "Comfort and moisture level"]
    ]

    for i, row in enumerate(table_data):
        for j, val in enumerate(row):
            cell = tbl.cell(i + 1, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = LIGHT_BLUE if i % 2 == 0 else WHITE
            cp = cell.text_frame.paragraphs[0]
            cp.text = val
            cp.font.name = "Arial"
            cp.font.size = Pt(13)
            cp.font.color.rgb = BLACK

    # Bottom notes about website
    wbox = s7.shapes.add_textbox(Inches(1.2), Inches(4.7), Inches(11.0), Inches(2.2))
    wtf = wbox.text_frame
    wtf.word_wrap = True

    web_notes = [
        "• The system uploads new data to ThingSpeak every 20 seconds.",
        "• The website automatically refreshes every 2 seconds to show live updates.",
        "• A clean 24-hour graph draws temperature and humidity history directly in the browser.",
        "• Live Website URL: https://smart-weather-station.vercel.app/",
        "• Project Code on GitHub: https://github.com/prasoon9999/smart-weather-station"
    ]

    for i, wn in enumerate(web_notes):
        p = wtf.paragraphs[0] if i == 0 else wtf.add_paragraph()
        p.text = wn
        p.font.name = "Arial"
        p.font.size = Pt(14)
        p.space_before = Pt(8)
        p.font.color.rgb = BLACK
        if "Live Website URL" in wn or "Project Code" in wn:
            p.font.bold = True

    # ==========================================
    # SLIDE 8: CONCLUSION (Matches PDF Page 7 - Easy Words)
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    add_common_header(s8, "CONCLUSION")

    box8 = s8.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(11.0), Inches(5.2))
    tf8 = box8.text_frame
    tf8.word_wrap = True

    concl_points = [
        "A low-cost and working Smart Weather Station was successfully built using ESP32.",
        "The DHT22 sensor gives accurate temperature and humidity readings in real time.",
        "The OLED display allows us to easily check values without needing an internet connection.",
        "The ThingSpeak cloud and Vercel website make it simple for anyone to see weather data from anywhere.",
        "Future Scope:",
        "    • Add a pressure sensor (BMP280) to measure atmospheric pressure and altitude.",
        "    • Add a small solar panel and battery so it can run outdoors without a power plug.",
        "    • Send SMS or mobile alerts if the temperature gets too high."
    ]

    for i, pt in enumerate(concl_points):
        p = tf8.paragraphs[0] if i == 0 else tf8.add_paragraph()
        p.text = pt if pt.startswith("Future") or pt.startswith("    ") else "•  " + pt
        p.font.name = "Arial"
        p.font.size = Pt(16 if not pt.startswith("Future") else 17)
        p.font.bold = pt.startswith("Future")
        p.space_before = Pt(8 if pt.startswith("    ") else 16)
        p.font.color.rgb = BLACK

    # ==========================================
    # SLIDE 9: REFERENCES:- (Matches PDF Page 8 - Exact Format)
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    add_common_header(s9, "REFERENCES:-")

    rows_ref = 6
    cols_ref = 2
    ref_table = s9.shapes.add_table(rows_ref, cols_ref, Inches(1.2), Inches(1.8), Inches(11.0), Inches(4.8))
    rtbl = ref_table.table
    rtbl.columns[0].width = Inches(1.4)
    rtbl.columns[1].width = Inches(9.6)

    # Header
    c0 = rtbl.cell(0, 0)
    c0.fill.solid()
    c0.fill.fore_color.rgb = BLUE_HEADER
    p_h0 = c0.text_frame.paragraphs[0]
    p_h0.text = "S. No."
    p_h0.font.name = "Arial"
    p_h0.font.size = Pt(14)
    p_h0.font.bold = True
    p_h0.font.color.rgb = WHITE
    p_h0.alignment = PP_ALIGN.CENTER

    c1 = rtbl.cell(0, 1)
    c1.fill.solid()
    c1.fill.fore_color.rgb = BLUE_HEADER
    p_h1 = c1.text_frame.paragraphs[0]
    p_h1.text = "Research Paper / Datasheet"
    p_h1.font.name = "Arial"
    p_h1.font.size = Pt(14)
    p_h1.font.bold = True
    p_h1.font.color.rgb = WHITE

    refs = [
        "\"IoT Based Weather Monitoring System Using ESP32 and Cloud Analytics\", IEEE Conference, 2021. {BASE PAPER}",
        "Espressif Systems, \"ESP32 Series Datasheet - Wi-Fi & Bluetooth Microcontroller\", 2023.",
        "Aosong Electronics, \"DHT22 (AM2302) Temperature and Humidity Sensor Datasheet\".",
        "Solomon Systech, \"SSD1306 0.96 inch OLED Display Driver Datasheet\", 2018.",
        "MathWorks, \"ThingSpeak IoT Open Cloud Platform Documentation\", 2024."
    ]

    for i, ref in enumerate(refs):
        cell_num = rtbl.cell(i + 1, 0)
        cell_num.fill.solid()
        cell_num.fill.fore_color.rgb = LIGHT_BLUE if i % 2 == 0 else WHITE
        cp0 = cell_num.text_frame.paragraphs[0]
        cp0.text = str(i + 1) + "."
        cp0.font.name = "Arial"
        cp0.font.size = Pt(13)
        cp0.alignment = PP_ALIGN.CENTER
        cp0.font.color.rgb = BLACK

        cell_txt = rtbl.cell(i + 1, 1)
        cell_txt.fill.solid()
        cell_txt.fill.fore_color.rgb = LIGHT_BLUE if i % 2 == 0 else WHITE
        cp1 = cell_txt.text_frame.paragraphs[0]
        cp1.text = ref
        cp1.font.name = "Arial"
        cp1.font.size = Pt(13)
        cp1.font.color.rgb = BLACK

    # ==========================================
    # SLIDE 10: THANK YOU (Matches PDF Page 9)
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    ty_box = s10.shapes.add_textbox(Inches(1.0), Inches(2.6), Inches(11.333), Inches(2.2))
    ty_tf = ty_box.text_frame
    ty_tf.word_wrap = True
    typ = ty_tf.paragraphs[0]
    typ.text = "THANK YOU"
    typ.font.name = "Arial"
    typ.font.size = Pt(56)
    typ.font.bold = True
    typ.font.color.rgb = RGBColor(120, 140, 160)
    typ.alignment = PP_ALIGN.CENTER

    out_path = r"c:\Users\dpras\Documents\smart weather station\Smart_Weather_Station_Presentation.pptx"
    prs.save(out_path)
    print(f"Presentation created in exact sample format at: {out_path}")

if __name__ == '__main__':
    create_simple_format_deck()
