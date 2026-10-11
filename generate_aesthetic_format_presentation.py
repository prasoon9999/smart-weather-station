import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_aesthetic_exact_format_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # --- ELEGANT PROFESSIONAL PALETTE ---
    # Academic & Clean: Navy Blue, Deep Slate, Subtle Accent Red & Blue
    NAVY_PRIMARY = RGBColor(26, 54, 93)      # #1A365D - Deep Royal Navy
    SLATE_DARK = RGBColor(30, 41, 59)        # #1E293B - Crisp Charcoal text
    TEXT_MUTED = RGBColor(100, 116, 139)     # #64748B - Refined Subtext
    ACCENT_RED = RGBColor(185, 28, 28)       # #B91C1C - Mentor Red
    ACCENT_TEAL = RGBColor(13, 148, 136)     # #0D9488 - Subtle Tech Accent
    BG_CARD = RGBColor(248, 250, 252)        # #F8FAFC - Soft crisp surface
    BORDER_LIGHT = RGBColor(226, 232, 240)   # #E2E8F0 - Thin subtle border
    WHITE = RGBColor(255, 255, 255)

    def add_page_decorations(slide, title_text):
        # Top subtle colored rule
        top_rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.08))
        top_rule.fill.solid()
        top_rule.fill.fore_color.rgb = NAVY_PRIMARY
        top_rule.line.fill.background()

        # Title text centered, exactly in the sample position but with elegant typography
        tbox = slide.shapes.add_textbox(Inches(1.0), Inches(0.55), Inches(11.333), Inches(0.85))
        tf = tbox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.name = "Arial"
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = NAVY_PRIMARY
        p.alignment = PP_ALIGN.CENTER

        # Clean aesthetic accent line below title
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5.666), Inches(1.45), Inches(2.0), Inches(0.04))
        line.fill.solid()
        line.fill.fore_color.rgb = ACCENT_TEAL
        line.line.fill.background()

    # ==========================================
    # SLIDE 1: EXACT TITLE SLIDE FORMAT (AESTHETIC)
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)

    # Top accent bar
    bar1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.1))
    bar1.fill.solid()
    bar1.fill.fore_color.rgb = NAVY_PRIMARY
    bar1.line.fill.background()

    # College Name & Dept Box (Centered, clean typography)
    top_box = s1.shapes.add_textbox(Inches(1.0), Inches(0.6), Inches(11.333), Inches(1.6))
    ttf = top_box.text_frame
    ttf.word_wrap = True
    
    p0 = ttf.paragraphs[0]
    p0.text = "Galgotias College of Engineering & Technology, Greater Noida"
    p0.font.name = "Arial"
    p0.font.size = Pt(24)
    p0.font.bold = True
    p0.font.color.rgb = NAVY_PRIMARY
    p0.alignment = PP_ALIGN.CENTER

    p1 = ttf.add_paragraph()
    p1.text = "DEPARTMENT OF ELECTRONIC AND COMMUNICATION ENGINEERING"
    p1.font.name = "Arial"
    p1.font.size = Pt(13)
    p1.font.bold = True
    p1.space_before = Pt(8)
    p1.font.color.rgb = SLATE_DARK
    p1.alignment = PP_ALIGN.CENTER

    # Elegant divider line
    div_line = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.5), Inches(2.1), Inches(4.333), Inches(0.03))
    div_line.fill.solid()
    div_line.fill.fore_color.rgb = ACCENT_TEAL
    div_line.line.fill.background()

    # Central Project Title Frame
    title_frame = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(2.6), Inches(10.333), Inches(2.0))
    title_frame.fill.solid()
    title_frame.fill.fore_color.rgb = BG_CARD
    title_frame.line.color.rgb = BORDER_LIGHT
    title_frame.line.width = Pt(1.5)

    c_tf = title_frame.text_frame
    c_tf.word_wrap = True
    c_tf.margin_top = Inches(0.35)

    cp = c_tf.paragraphs[0]
    cp.text = "Smart Weather Station"
    cp.font.name = "Arial"
    cp.font.size = Pt(36)
    cp.font.bold = True
    cp.font.color.rgb = NAVY_PRIMARY
    cp.alignment = PP_ALIGN.CENTER

    cp_sub = c_tf.add_paragraph()
    cp_sub.text = "Using ESP32, DHT22 Sensor and Cloud Dashboard"
    cp_sub.font.name = "Arial"
    cp_sub.font.size = Pt(18)
    cp_sub.space_before = Pt(8)
    cp_sub.font.color.rgb = ACCENT_TEAL
    cp_sub.alignment = PP_ALIGN.CENTER

    # Mentor (Bottom-Left - Matching format)
    m_box = s1.shapes.add_textbox(Inches(1.5), Inches(5.0), Inches(4.5), Inches(1.8))
    mtf = m_box.text_frame
    mp = mtf.paragraphs[0]
    mp.text = "MENTOR:"
    mp.font.name = "Arial"
    mp.font.size = Pt(16)
    mp.font.bold = True
    mp.font.color.rgb = ACCENT_RED

    mp_name = mtf.add_paragraph()
    mp_name.text = "[Mentor Name]"
    mp_name.font.name = "Arial"
    mp_name.font.size = Pt(15)
    mp_name.space_before = Pt(4)
    mp_name.font.bold = True
    mp_name.font.color.rgb = SLATE_DARK

    mp_dept = mtf.add_paragraph()
    mp_dept.text = "Assistant Professor, Dept of ECE"
    mp_dept.font.name = "Arial"
    mp_dept.font.size = Pt(13)
    mp_dept.font.color.rgb = TEXT_MUTED

    # Presented By (Bottom-Right - Matching format)
    p_box = s1.shapes.add_textbox(Inches(7.8), Inches(5.0), Inches(4.5), Inches(1.8))
    ptf = p_box.text_frame
    pp = ptf.paragraphs[0]
    pp.text = "PRESENTED BY:"
    pp.font.name = "Arial"
    pp.font.size = Pt(16)
    pp.font.bold = True
    pp.font.color.rgb = NAVY_PRIMARY

    students = [
        "1. Prasoon (Roll No: ____________)",
        "2. Name (Roll No: ____________)",
        "3. Name (Roll No: ____________)"
    ]
    for s in students:
        p_s = ptf.add_paragraph()
        p_s.text = s
        p_s.font.name = "Arial"
        p_s.font.size = Pt(14)
        p_s.space_before = Pt(4)
        p_s.font.color.rgb = SLATE_DARK

    # ==========================================
    # SLIDE 2: CONTENTS (Aesthetic Cards in 2 Columns)
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    add_page_decorations(s2, "CONTENTS")

    contents_list = [
        ("❖  Introduction", "Why weather monitoring matters and how IoT helps."),
        ("❖  Motivation", "Problems with old stations and the need for local data."),
        ("❖  Objectives", "What we wanted to build and achieve in this project."),
        ("❖  Description", "Hardware parts, circuit connections, and working flow."),
        ("❖  Conclusion", "What we achieved and ideas for future improvements."),
        ("❖  References", "Research papers, books, and component datasheets.")
    ]

    for idx, (title, sub) in enumerate(contents_list):
        row = idx // 2
        col = idx % 2
        x = Inches(1.4 + col * 5.4)
        y = Inches(1.9 + row * 1.55)

        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.1), Inches(1.3))
        card.fill.solid()
        card.fill.fore_color.rgb = BG_CARD
        card.line.color.rgb = BORDER_LIGHT
        card.line.width = Pt(1)

        ctf = card.text_frame
        ctf.word_wrap = True
        ctf.margin_left = Inches(0.3)
        ctf.margin_top = Inches(0.2)

        p = ctf.paragraphs[0]
        p.text = title
        p.font.name = "Arial"
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = NAVY_PRIMARY

        p_sub = ctf.add_paragraph()
        p_sub.text = sub
        p_sub.font.name = "Arial"
        p_sub.font.size = Pt(12)
        p_sub.space_before = Pt(4)
        p_sub.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 3: INTRODUCTION (Clean aesthetic cards)
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    add_page_decorations(s3, "INTRODUCTION")

    intro_points = [
        "Weather conditions like temperature and humidity affect our daily lives, health, and work.",
        "Traditional weather stations are big, expensive, and cannot be checked easily from home.",
        "Internet of Things (IoT) allows us to connect small sensors to the internet to collect weather data automatically.",
        "In this project, we built a low-cost smart weather station using an ESP32 board and a DHT22 sensor.",
        "The system reads live temperature and humidity, shows them on a small screen, and sends the data to a website so anyone can view it anytime."
    ]

    for i, pt in enumerate(intro_points):
        y = Inches(1.85 + i * 1.0)
        card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), y, Inches(10.933), Inches(0.85))
        card.fill.solid()
        card.fill.fore_color.rgb = BG_CARD
        card.line.color.rgb = BORDER_LIGHT
        card.line.width = Pt(1)

        ctf = card.text_frame
        ctf.word_wrap = True
        ctf.margin_left = Inches(0.3)
        ctf.margin_top = Inches(0.18)
        p = ctf.paragraphs[0]
        p.text = "•   " + pt
        p.font.name = "Arial"
        p.font.size = Pt(15)
        p.font.color.rgb = SLATE_DARK

    # ==========================================
    # SLIDE 4: MOTIVATION
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    add_page_decorations(s4, "MOTIVATION")

    mot_points = [
        ("Need for local weather data:", " City weather reports are taken far away (like airports) and do not show the exact weather in our room, office, or farm."),
        ("Low cost:", " Commercial weather devices cost a lot of money, but this project can be made at a very low cost using simple parts."),
        ("Real-time monitoring:", " Instead of writing down values by hand, the system updates automatically every few seconds."),
        ("Access from anywhere:", " Anyone with a phone or computer can open our website and check the live weather and past 24-hour charts.")
    ]

    for i, (bold_txt, regular_txt) in enumerate(mot_points):
        y = Inches(1.9 + i * 1.25)
        card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), y, Inches(10.933), Inches(1.1))
        card.fill.solid()
        card.fill.fore_color.rgb = BG_CARD
        card.line.color.rgb = BORDER_LIGHT
        card.line.width = Pt(1)

        ctf = card.text_frame
        ctf.word_wrap = True
        ctf.margin_left = Inches(0.3)
        ctf.margin_top = Inches(0.2)
        p = ctf.paragraphs[0]
        
        r1 = p.add_run()
        r1.text = "•   " + bold_txt
        r1.font.bold = True
        r1.font.size = Pt(16)
        r1.font.name = "Arial"
        r1.font.color.rgb = NAVY_PRIMARY

        r2 = p.add_run()
        r2.text = regular_txt
        r2.font.bold = False
        r2.font.size = Pt(15)
        r2.font.name = "Arial"
        r2.font.color.rgb = SLATE_DARK

    # ==========================================
    # SLIDE 5: OBJECTIVE:-
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    add_page_decorations(s5, "OBJECTIVE:-")

    objectives = [
        "To build a simple and working smart weather station using ESP32 microcontroller.",
        "To accurately measure room temperature and humidity using the DHT22 sensor.",
        "To display the current readings on a small OLED screen placed directly on the board.",
        "To calculate Heat Index (how hot it feels) and Dew Point to check human comfort.",
        "To send all sensor data to the ThingSpeak cloud platform over Wi-Fi every 20 seconds.",
        "To show live readings and a 24-hour history graph on an online website hosted on Vercel."
    ]

    for i, pt in enumerate(objectives):
        y = Inches(1.8 + i * 0.86)
        card = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), y, Inches(10.933), Inches(0.74))
        card.fill.solid()
        card.fill.fore_color.rgb = BG_CARD
        card.line.color.rgb = BORDER_LIGHT
        card.line.width = Pt(1)

        ctf = card.text_frame
        ctf.word_wrap = True
        ctf.margin_left = Inches(0.3)
        ctf.margin_top = Inches(0.14)
        p = ctf.paragraphs[0]
        p.text = "✔   " + pt
        p.font.name = "Arial"
        p.font.size = Pt(15)
        p.font.color.rgb = SLATE_DARK

    # ==========================================
    # SLIDE 6: DESCRIPTION - HARDWARE & CONNECTIONS
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    add_page_decorations(s6, "Description")

    # 3 Aesthetic Component Cards on Top
    comps = [
        ("ESP32 Microcontroller", "The main brain with built-in Wi-Fi to collect sensor data and send it to the cloud."),
        ("DHT22 Sensor", "Measures temperature (-40°C to 80°C) and humidity (0 to 100%) with high accuracy."),
        ("SSD1306 OLED Display", "Small 0.96-inch screen that shows live values directly on the device.")
    ]

    for i, (title, desc) in enumerate(comps):
        x = Inches(1.2 + i * 3.75)
        card = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.8), Inches(3.45), Inches(2.2))
        card.fill.solid()
        card.fill.fore_color.rgb = BG_CARD
        card.line.color.rgb = BORDER_LIGHT
        card.line.width = Pt(1)

        ctf = card.text_frame
        ctf.word_wrap = True
        ctf.margin_left = Inches(0.25)
        ctf.margin_top = Inches(0.2)

        p = ctf.paragraphs[0]
        p.text = f"{i+1}. {title}"
        p.font.name = "Arial"
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = NAVY_PRIMARY

        p2 = ctf.add_paragraph()
        p2.text = desc
        p2.font.name = "Arial"
        p2.font.size = Pt(13)
        p2.space_before = Pt(8)
        p2.font.color.rgb = SLATE_DARK

    # Bottom Full-width Connections Card
    bot_card = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(4.3), Inches(10.933), Inches(2.6))
    bot_card.fill.solid()
    bot_card.fill.fore_color.rgb = BG_CARD
    bot_card.line.color.rgb = ACCENT_TEAL
    bot_card.line.width = Pt(1.5)

    btf = bot_card.text_frame
    btf.word_wrap = True
    btf.margin_left = Inches(0.3)
    btf.margin_top = Inches(0.2)

    bp = btf.paragraphs[0]
    bp.text = "Wiring and Connections:"
    bp.font.name = "Arial"
    bp.font.size = Pt(16)
    bp.font.bold = True
    bp.font.color.rgb = NAVY_PRIMARY

    wlines = [
        "• DHT22 data pin is connected to ESP32 Pin GPIO 4.",
        "• OLED display is connected using I2C wires (SDA to GPIO 21, SCL to GPIO 22).",
        "• 3.3V and GND power pins are shared between the sensor and screen.",
        "Working Flow:",
        "    DHT22 reads data  ➔  ESP32 calculates values  ➔  Shows on OLED  ➔  Sends to ThingSpeak Cloud  ➔  Vercel Website"
    ]
    for wl in wlines:
        p = btf.add_paragraph()
        p.text = wl
        p.font.name = "Arial"
        p.font.size = Pt(14)
        p.space_before = Pt(4)
        if "Working Flow" in wl:
            p.font.bold = True
            p.font.color.rgb = NAVY_PRIMARY
        elif "DHT22 reads data" in wl:
            p.font.bold = True
            p.font.color.rgb = ACCENT_TEAL
        else:
            p.font.color.rgb = SLATE_DARK

    # ==========================================
    # SLIDE 7: DESCRIPTION - RESULTS & WEBSITE
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    add_page_decorations(s7, "Description")

    # Table on Left
    rows = 5
    cols = 3
    tbl_shape = s7.shapes.add_table(rows, cols, Inches(1.2), Inches(1.85), Inches(5.8), Inches(2.6))
    tbl = tbl_shape.table
    tbl.columns[0].width = Inches(2.0)
    tbl.columns[1].width = Inches(1.8)
    tbl.columns[2].width = Inches(2.0)

    headers = ["Parameter", "Live Value", "Meaning"]
    for j, h in enumerate(headers):
        c = tbl.cell(0, j)
        c.fill.solid()
        c.fill.fore_color.rgb = NAVY_PRIMARY
        cp = c.text_frame.paragraphs[0]
        cp.text = h
        cp.font.name = "Arial"
        cp.font.size = Pt(13)
        cp.font.bold = True
        cp.font.color.rgb = WHITE

    table_data = [
        ["Temperature", "26.8 °C", "Room temperature"],
        ["Humidity", "69.8 %", "Moisture in air"],
        ["Heat Index", "28.5 °C", "Real feel temperature"],
        ["Dew Point", "20.8 °C", "Moisture comfort"]
    ]
    for i, row in enumerate(table_data):
        for j, val in enumerate(row):
            c = tbl.cell(i + 1, j)
            c.fill.solid()
            c.fill.fore_color.rgb = BG_CARD if i % 2 == 0 else WHITE
            cp = c.text_frame.paragraphs[0]
            cp.text = val
            cp.font.name = "Arial"
            cp.font.size = Pt(13)
            cp.font.color.rgb = SLATE_DARK

    # Right Card: Live Website Features & Links
    rcard = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.3), Inches(1.85), Inches(4.833), Inches(4.8))
    rcard.fill.solid()
    rcard.fill.fore_color.rgb = BG_CARD
    rcard.line.color.rgb = ACCENT_TEAL
    rcard.line.width = Pt(1.5)

    rtf = rcard.text_frame
    rtf.word_wrap = True
    rtf.margin_left = Inches(0.3)
    rtf.margin_top = Inches(0.25)

    rp = rtf.paragraphs[0]
    rp.text = "Live Website & Cloud Features:"
    rp.font.name = "Arial"
    rp.font.size = Pt(16)
    rp.font.bold = True
    rp.font.color.rgb = NAVY_PRIMARY

    rpoints = [
        "• ESP32 sends new data to ThingSpeak every 20 seconds.",
        "• The website automatically refreshes every 2 seconds to show live updates.",
        "• A clean 24-hour graph draws temperature and humidity history directly in browser.",
        "• Live Website URL:\n   https://smart-weather-station.vercel.app/",
        "• Project Code on GitHub:\n   https://github.com/prasoon9999/smart-weather-station"
    ]
    for rpt in rpoints:
        p = rtf.add_paragraph()
        p.text = rpt
        p.font.name = "Arial"
        p.font.size = Pt(13)
        p.space_before = Pt(10)
        if "Live Website URL" in rpt or "Project Code" in rpt:
            p.font.bold = True
            p.font.color.rgb = NAVY_PRIMARY
        else:
            p.font.color.rgb = SLATE_DARK

    # Bottom note on Left
    l_note = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(4.7), Inches(5.8), Inches(1.95))
    l_note.fill.solid()
    l_note.fill.fore_color.rgb = BG_CARD
    l_note.line.color.rgb = BORDER_LIGHT
    l_note.line.width = Pt(1)
    ln_tf = l_note.text_frame
    ln_tf.margin_left = Inches(0.25)
    ln_tf.margin_top = Inches(0.2)
    ln_p = ln_tf.paragraphs[0]
    ln_p.text = "Telemetry Verification:"
    ln_p.font.name = "Arial"
    ln_p.font.size = Pt(15)
    ln_p.font.bold = True
    ln_p.font.color.rgb = NAVY_PRIMARY
    
    ln_p2 = ln_tf.add_paragraph()
    ln_p2.text = "All readings have been verified live on MathWorks ThingSpeak Channel 3531601 with over 60 continuous entries received from our hardware prototype."
    ln_p2.font.name = "Arial"
    ln_p2.font.size = Pt(13)
    ln_p2.space_before = Pt(6)
    ln_p2.font.color.rgb = SLATE_DARK

    # ==========================================
    # SLIDE 8: CONCLUSION
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    add_page_decorations(s8, "CONCLUSION")

    # Conclusion Card Top
    c_card = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(1.8), Inches(10.933), Inches(2.4))
    c_card.fill.solid()
    c_card.fill.fore_color.rgb = BG_CARD
    c_card.line.color.rgb = BORDER_LIGHT
    c_card.line.width = Pt(1)

    ctf = c_card.text_frame
    ctf.word_wrap = True
    ctf.margin_left = Inches(0.3)
    ctf.margin_top = Inches(0.2)

    cp = ctf.paragraphs[0]
    cp.text = "Summary of Work:"
    cp.font.name = "Arial"
    cp.font.size = Pt(16)
    cp.font.bold = True
    cp.font.color.rgb = NAVY_PRIMARY

    c_points = [
        "• A low-cost and working Smart Weather Station was successfully built using ESP32.",
        "• The DHT22 sensor gives accurate temperature and humidity readings in real time.",
        "• The OLED display allows us to easily check values without needing an internet connection.",
        "• The ThingSpeak cloud and Vercel website make it simple for anyone to see weather data from anywhere."
    ]
    for cpt in c_points:
        p = ctf.add_paragraph()
        p.text = cpt
        p.font.name = "Arial"
        p.font.size = Pt(14)
        p.space_before = Pt(4)
        p.font.color.rgb = SLATE_DARK

    # Future Scope Card Bottom
    f_card = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(4.45), Inches(10.933), Inches(2.4))
    f_card.fill.solid()
    f_card.fill.fore_color.rgb = BG_CARD
    f_card.line.color.rgb = ACCENT_TEAL
    f_card.line.width = Pt(1.5)

    ftf = f_card.text_frame
    ftf.word_wrap = True
    ftf.margin_left = Inches(0.3)
    ftf.margin_top = Inches(0.2)

    fp = ftf.paragraphs[0]
    fp.text = "Future Scope:"
    fp.font.name = "Arial"
    fp.font.size = Pt(16)
    fp.font.bold = True
    fp.font.color.rgb = ACCENT_TEAL

    f_points = [
        "• Add a pressure sensor (BMP280) to measure atmospheric pressure and altitude.",
        "• Add a small solar panel and battery so it can run outdoors without a power plug.",
        "• Send SMS or mobile alerts if the temperature gets too high."
    ]
    for fpt in f_points:
        p = ftf.add_paragraph()
        p.text = fpt
        p.font.name = "Arial"
        p.font.size = Pt(14)
        p.space_before = Pt(6)
        p.font.color.rgb = SLATE_DARK

    # ==========================================
    # SLIDE 9: REFERENCES:-
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    add_page_decorations(s9, "REFERENCES:-")

    ref_table_shape = s9.shapes.add_table(6, 2, Inches(1.2), Inches(1.85), Inches(10.933), Inches(4.8))
    rtbl = ref_table_shape.table
    rtbl.columns[0].width = Inches(1.4)
    rtbl.columns[1].width = Inches(9.533)

    c0 = rtbl.cell(0, 0)
    c0.fill.solid()
    c0.fill.fore_color.rgb = NAVY_PRIMARY
    p_h0 = c0.text_frame.paragraphs[0]
    p_h0.text = "S. No."
    p_h0.font.name = "Arial"
    p_h0.font.size = Pt(14)
    p_h0.font.bold = True
    p_h0.font.color.rgb = WHITE
    p_h0.alignment = PP_ALIGN.CENTER

    c1 = rtbl.cell(0, 1)
    c1.fill.solid()
    c1.fill.fore_color.rgb = NAVY_PRIMARY
    p_h1 = c1.text_frame.paragraphs[0]
    p_h1.text = "Research Paper / Datasheet"
    p_h1.font.name = "Arial"
    p_h1.font.size = Pt(14)
    p_h1.font.bold = True
    p_h1.font.color.rgb = WHITE

    refs = [
        "\"IoT Based Weather Monitoring System Using ESP32 and Cloud Analytics\", IEEE Conference, 2021. {BASE PAPER}",
        "Espressif Systems, \"ESP32 Series Datasheet - Wi-Fi & Bluetooth Microcontroller\", 2023.",
        "Aosong Electronics, \"DHT22 (AM2302) Temperature and Humidity Sensor Datasheet\".",
        "Solomon Systech, \"SSD1306 0.96 inch OLED Display Driver Datasheet\", 2018.",
        "MathWorks, \"ThingSpeak IoT Open Cloud Platform Documentation\", 2024."
    ]

    for i, ref in enumerate(refs):
        cell_num = rtbl.cell(i + 1, 0)
        cell_num.fill.solid()
        cell_num.fill.fore_color.rgb = BG_CARD if i % 2 == 0 else WHITE
        cp0 = cell_num.text_frame.paragraphs[0]
        cp0.text = str(i + 1) + "."
        cp0.font.name = "Arial"
        cp0.font.size = Pt(13)
        cp0.font.bold = True
        cp0.alignment = PP_ALIGN.CENTER
        cp0.font.color.rgb = SLATE_DARK

        cell_txt = rtbl.cell(i + 1, 1)
        cell_txt.fill.solid()
        cell_txt.fill.fore_color.rgb = BG_CARD if i % 2 == 0 else WHITE
        cp1 = cell_txt.text_frame.paragraphs[0]
        cp1.text = ref
        cp1.font.name = "Arial"
        cp1.font.size = Pt(13)
        cp1.font.color.rgb = SLATE_DARK

    # ==========================================
    # SLIDE 10: THANK YOU
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    
    # Elegant central card
    t_card = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.5), Inches(1.8), Inches(8.333), Inches(3.8))
    t_card.fill.solid()
    t_card.fill.fore_color.rgb = BG_CARD
    t_card.line.color.rgb = BORDER_LIGHT
    t_card.line.width = Pt(1.5)

    t_tf = t_card.text_frame
    t_tf.word_wrap = True
    t_tf.margin_top = Inches(0.8)

    tp = t_tf.paragraphs[0]
    tp.text = "THANK YOU"
    tp.font.name = "Arial"
    tp.font.size = Pt(56)
    tp.font.bold = True
    tp.font.color.rgb = NAVY_PRIMARY
    tp.alignment = PP_ALIGN.CENTER

    tp_sub = t_tf.add_paragraph()
    tp_sub.text = "Questions & Answers"
    tp_sub.font.name = "Arial"
    tp_sub.font.size = Pt(20)
    tp_sub.font.bold = True
    tp_sub.space_before = Pt(12)
    tp_sub.font.color.rgb = ACCENT_TEAL
    tp_sub.alignment = PP_ALIGN.CENTER

    out_path = r"c:\Users\dpras\Documents\smart weather station\Smart_Weather_Station_Presentation.pptx"
    prs.save(out_path)
    print(f"Aesthetic exact-format presentation created at: {out_path}")

if __name__ == '__main__':
    create_aesthetic_exact_format_deck()
